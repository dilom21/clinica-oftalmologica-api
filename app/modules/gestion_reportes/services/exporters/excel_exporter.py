from datetime import datetime, timezone
from io import BytesIO

def export_xlsx(title, columns, rows):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"
    sheet.append([title])
    sheet.append([datetime.now(timezone.utc).isoformat()])
    sheet.append([column["label"] for column in columns])
    for row in rows:
        sheet.append([row.get(column["key"]) for column in columns])
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(column["label"]) + 2, 14), 36)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
